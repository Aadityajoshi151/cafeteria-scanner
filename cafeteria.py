#!/usr/bin/env python3
"""
cafeteria scanner - prototype v2
Reader + Veg/Non-Veg/Complete buttons + buzzer + OLED + local storage + Excel export
+ email on Complete (settings in settings.ini).
"""

import configparser
import csv
import os
import queue
import smtplib
import sqlite3
import subprocess
import sys
import termios
import threading
import time
import tty
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path

import serial
from gpiozero import Button, Buzzer
from luma.core.interface.serial import spi
from luma.core.render import canvas
from luma.oled.device import ssd1306
from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import ImageFont

# ---------------- Settings (these will move to a config file later) ----------------
BASE_DIR = Path(__file__).resolve().parent
EMPLOYEES_CSV = BASE_DIR / "employees.csv"   # columns: card_id,emp_id,name
DB_PATH = BASE_DIR / "cafeteria.db"
OUTPUT_DIR = BASE_DIR / "output"
SETTINGS_INI = BASE_DIR / "settings.ini"     # email settings (read each time Complete is pressed)

SERIAL_PORT = "/dev/serial0"
PIN_VEG, PIN_NONVEG, PIN_COMPLETE, PIN_BUZZER = 17, 27, 22, 18
OLED_DC, OLED_RST = 24, 25

CHOICE_TIMEOUT = 10   # seconds to press Veg/Non-Veg after tapping a card
MESSAGE_TIME = 3      # seconds a result message stays on screen
REPEAT_IGNORE = 2     # ignore the same card for this long (module repeats while a card is held)
QUIT_KEY = "0"        # press this key in the terminal to stop the script
IDLE_REFRESH = 5      # seconds between idle-screen updates (date, counts, clock warning)
CLOCK_WARNING = True  # show "Clock not set" until the time is synced from the internet
                      # (set to False once an RTC module keeps the time without network)

FONT_PATH = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"

VEG, NONVEG, COMPLETE = "VEG", "NON-VEG", "COMPLETE"

events = queue.Queue()   # ("card", n), ("button", name) or ("email", (ok, text)); used by the main loop


# ---------------- Card reader ----------------
def parse_frame(buf):
    """RDM6300 frame body: 10 hex chars of data + 2 hex chars of checksum. Returns card number or None."""
    try:
        frame = buf.decode("ascii")
        if len(frame) != 12:
            return None
        data, checksum = frame[:10], int(frame[10:], 16)
        x = 0
        for i in range(0, 10, 2):
            x ^= int(data[i:i + 2], 16)
        return int(data[2:], 16) if x == checksum else None
    except (UnicodeDecodeError, ValueError):
        return None


def reader_thread():
    ser = serial.Serial(SERIAL_PORT, 9600, timeout=1)
    buf, last_id, last_time = b"", None, 0.0
    while True:
        b = ser.read(1)
        if not b:
            continue
        if b == b"\x02":                 # start of frame
            buf = b""
        elif b == b"\x03":               # end of frame
            n = parse_frame(buf)
            buf = b""
            if n is None:
                continue
            now = time.monotonic()
            if n == last_id and now - last_time < REPEAT_IGNORE:
                last_time = now          # card still being held: ignore
                continue
            last_id, last_time = n, now
            events.put(("card", n))
        else:
            buf += b
            if len(buf) > 20:            # junk guard
                buf = b""


def card_formats(n):
    """All the ways a card number may be written in the employee list."""
    return {f"{n:010d}", str(n), f"{(n >> 16) & 0xFF:03d},{n & 0xFFFF:05d}"}


# ---------------- Employee list ----------------
class Employees:
    """Reads employees.csv and reloads it automatically whenever the file changes."""

    def __init__(self, path):
        self.path, self.mtime, self.by_card = path, None, {}

    def _reload_if_changed(self):
        try:
            mtime = self.path.stat().st_mtime
        except FileNotFoundError:
            if self.by_card:
                print("employees.csv not found - all cards will be unmatched")
            self.by_card, self.mtime = {}, None
            return
        if mtime == self.mtime:
            return
        self.mtime = mtime
        table = {}
        with open(self.path, newline="", encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                card_id = (row.get("card_id") or "").strip()
                if card_id:
                    table[card_id] = ((row.get("emp_id") or "").strip(),
                                      (row.get("name") or "").strip())
        self.by_card = table
        print(f"Loaded {len(table)} employees")

    def lookup(self, n):
        self._reload_if_changed()
        for key in card_formats(n):
            if key in self.by_card:
                return self.by_card[key]
        return None


# ---------------- Storage ----------------
def open_db():
    db = sqlite3.connect(DB_PATH)
    db.execute("PRAGMA synchronous=FULL")    # every scan is safely on disk before we beep OK
    db.execute("""CREATE TABLE IF NOT EXISTS lunch(
                    day TEXT, card_id TEXT, emp_id TEXT, name TEXT, meal TEXT, time TEXT,
                    PRIMARY KEY (day, card_id))""")
    db.execute("""CREATE TABLE IF NOT EXISTS log(
                    ts TEXT, card_id TEXT, meal TEXT, result TEXT)""")
    db.commit()
    return db


def record(db, n, emp, meal):
    """One row per card per day. Returns 'new', 'same' (already recorded) or 'changed'."""
    now = datetime.now()
    day, t = now.strftime("%Y-%m-%d"), now.strftime("%H:%M:%S")
    card_id = f"{n:010d}"
    emp_id, name = emp if emp else ("", "")
    row = db.execute("SELECT meal FROM lunch WHERE day=? AND card_id=?", (day, card_id)).fetchone()
    result = "new" if row is None else ("same" if row[0] == meal else "changed")
    if result != "same":
        db.execute("INSERT OR REPLACE INTO lunch VALUES (?,?,?,?,?,?)",
                   (day, card_id, emp_id, name, meal, t))
    db.execute("INSERT INTO log VALUES (?,?,?,?)",
               (now.isoformat(timespec="seconds"), card_id, meal, result))
    db.commit()
    return result


def today_counts(db):
    day = datetime.now().strftime("%Y-%m-%d")
    rows = dict(db.execute("SELECT meal, COUNT(*) FROM lunch WHERE day=? GROUP BY meal",
                           (day,)).fetchall())
    return rows.get(VEG, 0), rows.get(NONVEG, 0)


def export_today(db):
    """Writes output/lunch_YYYY-MM-DD.xlsx with a 'Lunch' sheet and an 'Unmatched' sheet."""
    day = datetime.now().strftime("%Y-%m-%d")
    rows = db.execute("SELECT emp_id, name, meal, time, card_id FROM lunch WHERE day=? ORDER BY time",
                      (day,)).fetchall()
    matched = [r for r in rows if r[0]]
    unmatched = [r for r in rows if not r[0]]
    bold = Font(bold=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Lunch"
    ws.append(["Emp ID", "Name", "Veg", "Non-Veg", "Time"])
    for emp_id, name, meal, t, _ in matched:
        ws.append([emp_id, name, "✓" if meal == VEG else "", "✓" if meal == NONVEG else "", t])
    ws.append([])
    ws.append(["", "Total",
               sum(1 for r in matched if r[2] == VEG),
               sum(1 for r in matched if r[2] == NONVEG), ""])
    for cell in ws[1] + ws[ws.max_row]:
        cell.font = bold
    for col, width in zip("ABCDE", (12, 28, 8, 10, 10)):
        ws.column_dimensions[col].width = width

    ws2 = wb.create_sheet("Unmatched")
    ws2.append(["Card ID", "Wiegand", "Meal", "Time", "Who (fill from register)"])
    for _, _, meal, t, card_id in unmatched:
        n = int(card_id)
        ws2.append([card_id, f"{(n >> 16) & 0xFF:03d},{n & 0xFFFF:05d}", meal, t, ""])
    for cell in ws2[1]:
        cell.font = bold
    for col, width in zip("ABCDE", (14, 12, 10, 10, 26)):
        ws2.column_dimensions[col].width = width

    OUTPUT_DIR.mkdir(exist_ok=True)
    path = OUTPUT_DIR / f"lunch_{day}.xlsx"
    wb.save(path)
    return path, len(rows), len(unmatched)


# ---------------- Email ----------------
class SettingsError(Exception):
    pass


def load_email_settings():
    cfg = configparser.ConfigParser(interpolation=None)   # no interpolation: '%' allowed in passwords
    if not cfg.read(SETTINGS_INI, encoding="utf-8") or "email" not in cfg:
        raise SettingsError("settings.ini missing or has no [email] section")
    e = cfg["email"]
    s = {
        "server": e.get("smtp_server", "").strip(),
        "port": e.getint("smtp_port", fallback=587),
        "username": e.get("username", "").strip(),
        "password": e.get("password", "").replace(" ", ""),
        "sender": e.get("sender", "").strip() or e.get("username", "").strip(),
        "to": [a.strip() for a in e.get("to", "").split(",") if a.strip()],
        "cc": [a.strip() for a in e.get("cc", "").split(",") if a.strip()],
        "subject": e.get("subject", "Canteen lunch - {date}"),
    }
    if not s["server"] or not s["sender"] or not s["to"]:
        raise SettingsError("smtp_server, username/sender and to are required")
    return s


def send_report(path, summary):
    """Emails the Excel file. Raises on any failure."""
    s = load_email_settings()
    day = datetime.now().strftime("%d %b %Y")
    msg = EmailMessage()
    msg["Subject"] = s["subject"].replace("{date}", day)
    msg["From"] = s["sender"]
    msg["To"] = ", ".join(s["to"])
    if s["cc"]:
        msg["Cc"] = ", ".join(s["cc"])
    msg.set_content(f"Canteen lunch entries for {day}.\n\n{summary}\n\n"
                    "The Excel file is attached. Cards not in the employee list are on the "
                    "'Unmatched' sheet.\n\n- Canteen scanner")
    msg.add_attachment(path.read_bytes(), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       filename=path.name)

    if s["port"] == 465:                      # SSL from the start
        smtp = smtplib.SMTP_SSL(s["server"], 465, timeout=30)
    else:                                     # usually 587: plain, then upgrade with STARTTLS
        smtp = smtplib.SMTP(s["server"], s["port"], timeout=30)
    with smtp:
        smtp.ehlo()
        if s["port"] != 465 and smtp.has_extn("starttls"):
            smtp.starttls()
            smtp.ehlo()
        if s["username"]:                     # some office mail relays need no login
            smtp.login(s["username"], s["password"])
        smtp.send_message(msg)


def email_worker(path, summary):
    """Runs in the background so scanning and the screen keep working while sending."""
    try:
        send_report(path, summary)
        result = (True, "")
    except SettingsError as e:
        print(f"Email settings problem: {e}")
        result = (False, "Check settings")
    except smtplib.SMTPAuthenticationError:
        result = (False, "Login failed")
    except smtplib.SMTPRecipientsRefused:
        result = (False, "Bad recipient")
    except smtplib.SMTPException as e:
        print(f"Mail server error: {e}")
        result = (False, "Mail server error")
    except OSError as e:                      # DNS failure, timeout, no Wi-Fi, blocked port
        print(f"Network error: {e}")
        result = (False, "No network")
    except Exception as e:
        print(f"Unexpected email error: {e!r}")
        result = (False, "Unknown error")
    events.put(("email", result))


# ---------------- Display and buzzer ----------------
def load_font(size):
    try:
        return ImageFont.truetype(FONT_PATH, size)
    except OSError:
        return ImageFont.load_default()


FONT_BIG, FONT_SMALL = load_font(16), load_font(11)


def show(device, big, small1="", small2=""):
    with canvas(device) as draw:
        draw.text((0, 0), big[:12], font=FONT_BIG, fill="white")
        draw.text((0, 26), small1[:20], font=FONT_SMALL, fill="white")
        draw.text((0, 44), small2[:20], font=FONT_SMALL, fill="white")


BEEPS = {                 # (on seconds, off seconds, count)
    "read": (0.05, 0.05, 1),
    "ok": (0.2, 0.1, 1),
    "twice": (0.08, 0.08, 2),
    "error": (0.1, 0.1, 3),
    "timeout": (0.6, 0.1, 1),
}


def beep(buzzer, kind):
    on, off, n = BEEPS[kind]
    buzzer.beep(on_time=on, off_time=off, n=n, background=True)


# ---------------- Clock ----------------
_clock = {"synced": False, "checked": 0.0}


def clock_synced():
    """True once the system clock has been set from the internet (checked at most every 30 s)."""
    if not CLOCK_WARNING:
        return True
    now = time.monotonic()
    if _clock["synced"] or now - _clock["checked"] < 30:
        return _clock["synced"]
    _clock["checked"] = now
    synced = Path("/run/systemd/timesync/synchronized").exists()   # created by systemd-timesyncd
    if not synced:
        try:
            out = subprocess.run(["timedatectl", "show", "-p", "NTPSynchronized", "--value"],
                                 capture_output=True, text=True, timeout=5).stdout
            synced = out.strip() == "yes"
        except (OSError, subprocess.SubprocessError):
            pass
    if synced:
        print(f"Clock synced: {datetime.now():%d %b %Y %H:%M}")
    _clock["synced"] = synced
    return synced


# ---------------- Quit key and shutdown ----------------
RESOURCES = {}   # things to close on exit, filled in by main()


def cleanup():
    """Clear the screen, silence the buzzer, release pins and files, restore the terminal."""
    print("\nStopping...")
    steps = [
        lambda: termios.tcsetattr(sys.stdin.fileno(), termios.TCSADRAIN, RESOURCES["terminal"])
        if RESOURCES.get("terminal") else None,
        lambda: RESOURCES["buzzer"].off(),
        lambda: [b.close() for b in RESOURCES["buttons"]],
        lambda: RESOURCES["buzzer"].close(),
        lambda: RESOURCES["device"].cleanup(),      # blanks the OLED
        lambda: RESOURCES["db"].close(),
    ]
    for step in steps:
        try:
            step()
        except Exception:
            pass                                # keep going: exiting matters more than tidiness
    print("Stopped.")

def keyboard_thread():
    """Reads single key presses (no Enter needed); QUIT_KEY stops the script."""
    while True:
        ch = sys.stdin.read(1)
        if ch == QUIT_KEY or ch == "":    # "" means the terminal closed
            events.put(("quit", None))
            return


def start_quit_key():
    """Only when run by hand in a terminal. Returns the old terminal settings to restore on exit."""
    if not sys.stdin.isatty():
        return None
    fd = sys.stdin.fileno()
    old = termios.tcgetattr(fd)
    tty.setcbreak(fd)                     # single keys, no Enter; Ctrl+C still works
    threading.Thread(target=keyboard_thread, daemon=True).start()
    return old


# ---------------- Main loop ----------------
def main():
    device = ssd1306(spi(port=0, device=0, gpio_DC=OLED_DC, gpio_RST=OLED_RST))
    buzzer = Buzzer(PIN_BUZZER)
    db = open_db()
    employees = Employees(EMPLOYEES_CSV)

    buttons = {
        VEG: Button(PIN_VEG, bounce_time=0.05),
        NONVEG: Button(PIN_NONVEG, bounce_time=0.05),
        COMPLETE: Button(PIN_COMPLETE, bounce_time=0.05),
    }
    for name, btn in buttons.items():
        btn.when_pressed = lambda n=name: events.put(("button", n))
    RESOURCES.update(device=device, buzzer=buzzer, db=db, buttons=list(buttons.values()))

    threading.Thread(target=reader_thread, daemon=True).start()

    pending = None        # (card number, employee or None, deadline) while waiting for a button
    message_until = 0.0   # a result message stays on screen until this time
    sending = False       # True while an email is being sent in the background
    deferred = None       # email result to show once the current customer is done

    idle = {"shown": None, "at": 0.0}

    def idle_screen(force=True):
        """Draws the idle screen. Redraws only if something changed, unless forced."""
        idle["at"] = time.monotonic()
        veg, nonveg = today_counts(db)
        bottom = datetime.now().strftime("%d %b %Y") if clock_synced() else "! Clock not set"
        lines = ("Tap your card", f"Veg {veg}  Non-Veg {nonveg}", bottom)
        if force or lines != idle["shown"]:
            show(device, *lines)
            idle["shown"] = lines

    def message(big, s1="", s2=""):
        nonlocal message_until
        show(device, big, s1, s2)
        message_until = time.monotonic() + MESSAGE_TIME

    idle_screen()
    beep(buzzer, "ok")
    RESOURCES["terminal"] = start_quit_key()
    print(f"Ready. Press {QUIT_KEY} to quit.")

    while True:
        try:
            kind, value = events.get(timeout=0.2)
        except queue.Empty:
            kind, value = None, None
        now = time.monotonic()

        if kind == "quit":
            return

        if kind == "card":
            emp = employees.lookup(value)
            pending = (value, emp, now + CHOICE_TIMEOUT)   # a new tap replaces any unfinished one
            message_until = 0.0
            beep(buzzer, "read")
            if emp:
                show(device, emp[1] or emp[0], "Press VEG or", "NON-VEG")
            else:
                show(device, "Unknown card", f"ID {value:010d}", "Press VEG/NON-VEG")
            print(f"Card {value:010d} -> {emp if emp else 'UNKNOWN'}")

        elif kind == "button" and value in (VEG, NONVEG):
            if pending is None:
                message("Tap card", "first, then press", "VEG or NON-VEG")
                beep(buzzer, "error")
            else:
                n, emp, _ = pending
                pending = None
                result = record(db, n, emp, value)
                who = (emp[1] or emp[0]) if emp else f"{n:010d}"
                if result == "same":
                    message(who, f"Already: {value}")
                    beep(buzzer, "twice")
                elif result == "changed":
                    message(who, f"Changed to {value}")
                    beep(buzzer, "ok")
                else:
                    message(who, f"{value}  OK")
                    beep(buzzer, "ok")
                print(f"{who}: {value} ({result})")

        elif kind == "button" and value == COMPLETE:
            if sending:
                message("Please wait", "Email still sending")
                beep(buzzer, "error")
            else:
                pending = None
                path, total, unmatched = export_today(db)
                veg, nonveg = today_counts(db)
                summary = (f"Veg: {veg}\nNon-Veg: {nonveg}\nTotal: {total}\n"
                           f"Unmatched cards (included above): {unmatched}")
                print(f"Exported {total} entries ({unmatched} unmatched) to {path}; emailing...")
                sending = True
                show(device, "Sending...", f"{total} entries", "Please wait")
                message_until = time.monotonic() + 120   # held until the email result arrives
                beep(buzzer, "read")
                threading.Thread(target=email_worker, args=(path, summary), daemon=True).start()

        elif kind == "email":
            sending = False
            ok, reason = value
            print("Email sent" if ok else f"Email FAILED: {reason}")
            screen = ("Email sent", "File delivered", "") if ok else \
                     ("Email FAILED", reason, "Press COMPLETE again")
            beep(buzzer, "twice" if ok else "error")
            if pending is None:
                message(*screen)
            else:
                deferred = screen                  # don't hide someone's Veg/Non-Veg prompt

        if pending and now > pending[2]:
            pending = None
            message("Timed out", "Tap card again")
            beep(buzzer, "timeout")
            print("Timed out waiting for a button")

        if pending is None and message_until and now > message_until:
            if deferred:
                message(*deferred)
                deferred = None
            else:
                message_until = 0.0
                idle_screen()

        # Keep the idle screen current: date changes, clock syncing after boot, midnight.
        if pending is None and not message_until and now - idle["at"] >= IDLE_REFRESH:
            idle_screen(force=False)


if __name__ == "__main__":
    exit_code = 0
    try:
        main()
    except KeyboardInterrupt:
        pass
    except Exception:
        import traceback
        traceback.print_exc()                   # show the error before exiting
        exit_code = 1
    finally:
        cleanup()
        # Background threads (card reader, key reader, email) can block a normal exit,
        # so leave immediately. Every scan is already saved to the database.
        sys.stdout.flush()
        os._exit(exit_code)